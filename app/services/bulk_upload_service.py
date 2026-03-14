import pandas as pd
from io import BytesIO
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session

from app.models.taxonomy import (
    Category,
    SubCategory,
    SubSubCategory,
    ProductType,
    Material,
    Color,
    Size,
)
from app.services.product_service import create_product


REQUIRED_COLUMNS = [
    "category",
    "subcategory",
    "product_type",
    "material",
]

OPTIONAL_COLUMNS = [
    "subsubcategory",
    "color",
    "size",
    "gtin",
    "initial_quantity",
    "title",
]

ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


def _build_lookup(db: Session, model, key_attr="short_code"):
    """Build a case-insensitive short_code → model instance lookup dict."""
    rows = db.query(model).all()
    return {getattr(r, key_attr).strip().upper(): r for r in rows}


def process_bulk_upload(
    db: Session,
    organization_id,
    file_bytes: bytes,
    filename: str,
):
    # Read Excel — skip row 2 (instructions row from template)
    try:
        df = pd.read_excel(BytesIO(file_bytes), engine="openpyxl", dtype=str, skiprows=[1])
    except Exception as e:
        return {
            "total_rows": 0,
            "created": 0,
            "failed": 0,
            "errors": [{"row": 0, "error": f"Cannot read Excel file: {e}"}],
        }

    df.columns = [c.strip().lower().replace(" ", "_") for c in df.columns]

    # Validate required columns exist
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        return {
            "total_rows": 0,
            "created": 0,
            "failed": 0,
            "errors": [{"row": 0, "error": f"Missing required columns: {', '.join(missing)}"}],
        }

    # Pre-load taxonomy lookups
    categories = _build_lookup(db, Category)
    subcategories = _build_lookup(db, SubCategory)
    subsubcategories = _build_lookup(db, SubSubCategory)
    product_types = _build_lookup(db, ProductType)
    materials = _build_lookup(db, Material)
    colors = _build_lookup(db, Color)
    sizes = _build_lookup(db, Size)

    created = 0
    errors = []

    for idx, row in df.iterrows():
        row_num = idx + 3  # Excel row (row 1=header, row 2=instructions, data starts row 3)

        try:
            # --- Resolve required taxonomy ---
            cat_code = str(row.get("category", "") or "").strip().upper()
            sub_code = str(row.get("subcategory", "") or "").strip().upper()
            pt_code = str(row.get("product_type", "") or "").strip().upper()
            mat_code = str(row.get("material", "") or "").strip().upper()

            category = categories.get(cat_code)
            if not category:
                errors.append({"row": row_num, "error": f"Unknown category: '{cat_code}'"})
                continue

            subcategory = subcategories.get(sub_code)
            if not subcategory:
                errors.append({"row": row_num, "error": f"Unknown subcategory: '{sub_code}'"})
                continue

            product_type = product_types.get(pt_code)
            if not product_type:
                errors.append({"row": row_num, "error": f"Unknown product_type: '{pt_code}'"})
                continue

            material = materials.get(mat_code)
            if not material:
                errors.append({"row": row_num, "error": f"Unknown material: '{mat_code}'"})
                continue

            # --- Resolve optional taxonomy ---
            subsub_code = str(row.get("subsubcategory", "") or "").strip().upper()
            subsubcategory = subsubcategories.get(subsub_code) if subsub_code else None

            color_code = str(row.get("color", "") or "").strip().upper()
            color = colors.get(color_code) if color_code else None

            size_code = str(row.get("size", "") or "").strip().upper()
            size = sizes.get(size_code) if size_code else None

            # --- Optional fields ---
            gtin = str(row.get("gtin", "") or "").strip() or None

            raw_qty = str(row.get("initial_quantity", "") or "").strip()
            initial_quantity = int(raw_qty) if raw_qty else 0

            title = str(row.get("title", "") or "").strip() or None

            # --- Create product ---
            product = create_product(
                db=db,
                organization_id=organization_id,
                category=category,
                subcategory=subcategory,
                subsubcategory=subsubcategory,
                product_type=product_type,
                material=material,
                color=color,
                size=size,
                gtin=gtin,
                initial_quantity=initial_quantity,
            )

            if title:
                product.title = title

            db.flush()
            created += 1

        except ValueError as e:
            errors.append({"row": row_num, "error": str(e)})
        except Exception as e:
            errors.append({"row": row_num, "error": str(e)})

    if created > 0:
        db.commit()

    return {
        "total_rows": len(df),
        "created": created,
        "failed": len(errors),
        "errors": errors,
    }


def _get_short_codes(db: Session, model) -> list[str]:
    """Return sorted list of short_codes for a taxonomy model."""
    rows = db.query(model.short_code).order_by(model.short_code).all()
    return [r[0] for r in rows]


def _get_code_name_pairs(db: Session, model) -> list[tuple[str, str]]:
    """Return sorted list of (short_code, name) tuples for a taxonomy model."""
    rows = db.query(model.short_code, model.name).order_by(model.short_code).all()
    return [(r[0], r[1]) for r in rows]


# Column descriptions for the instructions row
COLUMN_DESCRIPTIONS = {
    "category": "Required — Category short code (see Taxonomy Reference sheet)",
    "subcategory": "Required — Subcategory short code",
    "product_type": "Required — Product type short code",
    "material": "Required — Material short code",
    "subsubcategory": "Optional — Sub-subcategory short code",
    "color": "Optional — Color short code",
    "size": "Optional — Size short code",
    "gtin": "Optional — GTIN / UPC barcode",
    "initial_quantity": "Optional — Starting inventory (default: 0)",
    "title": "Optional — Product title",
}


def generate_bulk_template(db: Session) -> bytes:
    """Generate an Excel template with dropdown validations for taxonomy columns."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"

    # --- Styling ---
    header_font = Font(bold=True, color="FFFFFF", size=11)
    required_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    optional_fill = PatternFill(start_color="808080", end_color="808080", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        bottom=Side(style="thin", color="000000"),
    )
    hint_font = Font(italic=True, color="666666", size=10)
    hint_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    # --- Row 1: Headers ---
    for col_idx, col_name in enumerate(ALL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = required_fill if col_name in REQUIRED_COLUMNS else optional_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = 22

    # --- Row 2: Instructions / usage hints ---
    for col_idx, col_name in enumerate(ALL_COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=COLUMN_DESCRIPTIONS.get(col_name, ""))
        cell.font = hint_font
        cell.fill = hint_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.row_dimensions[2].height = 40

    # --- Hidden _Lookups sheet for dropdown values ---
    ref_ws = wb.create_sheet("_Lookups")
    ref_ws.sheet_state = "hidden"

    taxonomy_map = {
        "category": (Category, "A"),
        "subcategory": (SubCategory, "B"),
        "subsubcategory": (SubSubCategory, "C"),
        "product_type": (ProductType, "D"),
        "material": (Material, "E"),
        "color": (Color, "F"),
        "size": (Size, "G"),
    }

    col_index_map = {name: idx for idx, name in enumerate(ALL_COLUMNS, start=1)}
    max_data_rows = 500  # rows available for user to fill
    data_start_row = 3   # data starts after header + instructions

    for col_name, (model, ref_col) in taxonomy_map.items():
        codes = _get_short_codes(db, model)
        if not codes:
            continue

        # Write values to hidden sheet
        for row_idx, code in enumerate(codes, start=1):
            ref_ws.cell(row=row_idx, column=ord(ref_col) - ord("A") + 1, value=code)

        # Create data validation referencing the hidden sheet
        formula = f"_Lookups!${ref_col}$1:${ref_col}${len(codes)}"
        dv = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=(col_name not in REQUIRED_COLUMNS),
        )
        dv.error = f"Please select a valid {col_name}"
        dv.errorTitle = "Invalid Value"
        dv.prompt = f"Select {col_name}"
        dv.promptTitle = col_name.replace("_", " ").title()

        target_col_letter = ws.cell(row=1, column=col_index_map[col_name]).column_letter
        dv.add(f"{target_col_letter}{data_start_row}:{target_col_letter}{max_data_rows + data_start_row - 1}")
        ws.add_data_validation(dv)

    # Freeze rows 1-2 (header + instructions)
    ws.freeze_panes = "A3"

    # --- Taxonomy Reference sheet ---
    tax_ws = wb.create_sheet("Taxonomy Reference")

    ref_header_font = Font(bold=True, color="FFFFFF", size=11)
    ref_header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    section_font = Font(bold=True, size=12, color="2F5496")

    taxonomy_sections = [
        ("Category", Category),
        ("Subcategory", SubCategory),
        ("Sub-subcategory", SubSubCategory),
        ("Product Type", ProductType),
        ("Material", Material),
        ("Color", Color),
        ("Size", Size),
    ]

    current_row = 1

    for section_name, model in taxonomy_sections:
        pairs = _get_code_name_pairs(db, model)
        if not pairs:
            continue

        # Section title
        cell = tax_ws.cell(row=current_row, column=1, value=section_name)
        cell.font = section_font
        current_row += 1

        # Table header
        for col, label in [(1, "Short Code"), (2, "Name")]:
            cell = tax_ws.cell(row=current_row, column=col, value=label)
            cell.font = ref_header_font
            cell.fill = ref_header_fill
            cell.alignment = Alignment(horizontal="center")
        current_row += 1

        # Data rows
        for code, name in pairs:
            tax_ws.cell(row=current_row, column=1, value=code)
            tax_ws.cell(row=current_row, column=2, value=name)
            current_row += 1

        # Blank row between sections
        current_row += 1

    tax_ws.column_dimensions["A"].width = 18
    tax_ws.column_dimensions["B"].width = 35

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
